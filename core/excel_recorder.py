"""测试记录 Excel 导出模块。

每次测试结束后，按 test_id 从 SQLite 数据库读取主记录与测试项判定结果，
生成格式化 Excel 文件，按 测试记录/项目号/年月日/ 目录归档。
"""

import datetime
import os
import sqlite3
import sys
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _default_root_dir() -> str:
    """与 DBManager 保持一致：打包后锚定到 exe 目录，源码运行锚定到项目根目录"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def extract_project_no(recipe_name: str) -> str:
    """从配方名取项目号：按下划线取第一段，如 DJ2513_Aging66666 -> DJ2513"""
    name = (recipe_name or "").strip()
    if not name:
        return "未知项目"
    first = name.split("_")[0].strip()
    return first or "未知项目"


def _sanitize(text: str) -> str:
    text = "".join(c for c in (text or "") if c.isalnum() or c in "_-")
    return text or "UNKNOWN"


def _parse_time(value: str) -> Optional[datetime.datetime]:
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f"):
        try:
            return datetime.datetime.strptime(value or "", fmt)
        except ValueError:
            continue
    return None


def export_test_record(db_path: str, test_id: int, base_dir: Optional[str] = None) -> Optional[str]:
    """导出指定 test_id 的测试记录为 Excel，返回文件绝对路径；记录不存在返回 None。"""
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        main = conn.execute(
            "SELECT * FROM test_main WHERE test_id = ?", (test_id,)
        ).fetchone()
        if main is None:
            return None
        items = conn.execute(
            "SELECT * FROM test_items_results WHERE test_id = ? ORDER BY id ASC",
            (test_id,),
        ).fetchall()
    finally:
        conn.close()

    recipe_name = main["recipe_name"] or ""
    project_no = extract_project_no(recipe_name)
    start_dt = _parse_time(main["start_time"]) or datetime.datetime.now()
    date_str = start_dt.strftime("%Y-%m-%d")
    ts_str = start_dt.strftime("%Y%m%d_%H%M%S")
    result = main["result"] or ""
    sn = _sanitize(main["master_code"])

    if base_dir is None:
        base_dir = os.path.join(_default_root_dir(), "测试记录")
    out_dir = os.path.join(base_dir, project_no, date_str)
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"{sn}_{ts_str}_{result}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "测试记录"

    title_font = Font(size=14, bold=True)
    label_font = Font(bold=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    pass_font = Font(bold=True, color="006100")
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_font = Font(bold=True, color="9C0006")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    # 标题
    ws.merge_cells("A1:I1")
    cell = ws.cell(row=1, column=1, value="BMS 测试记录报告")
    cell.font = title_font
    cell.alignment = center
    ws.row_dimensions[1].height = 28

    # 主记录信息块（两列 标签/值 对）
    info_pairs = [
        ("项目号", project_no, "配方名称", recipe_name),
        ("产品条码(SN)", main["master_code"] or "", "开始时间", main["start_time"] or ""),
        ("结束时间", main["end_time"] or ""),
    ]
    row = 3
    for pair in info_pairs:
        label_a, value_a = pair[0], pair[1]
        ws.cell(row=row, column=1, value=label_a).font = label_font
        ws.cell(row=row, column=2, value=value_a).alignment = left
        if len(pair) > 2:
            ws.cell(row=row, column=4, value=pair[2]).font = label_font
            ws.cell(row=row, column=5, value=pair[3]).alignment = left
        row += 1

    ws.cell(row=row, column=1, value="总结果").font = label_font
    result_cell = ws.cell(row=row, column=2, value=result)
    result_cell.alignment = center
    if result == "PASS":
        result_cell.font = pass_font
        result_cell.fill = pass_fill
    else:
        result_cell.font = fail_font
        result_cell.fill = fail_fill
    row += 2

    # 测试项判定表
    headers = ["序号", "测试项", "下限", "上限", "实测值", "单位", "判定", "耗时(s)", "记录时间"]
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
    header_row = row
    row += 1

    for idx, item in enumerate(items, start=1):
        unit = item["unit"]
        if unit == "NULL":
            unit = ""
        values = [
            idx,
            item["item_name"],
            item["lower_limit"],
            item["upper_limit"],
            item["measured_value"],
            unit,
            item["result"],
            item["duration"] if "duration" in item.keys() else None,
            item["timestamp"],
        ]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=value)
            c.alignment = center if col != 2 else left
        verdict = (item["result"] or "").upper()
        verdict_cell = ws.cell(row=row, column=7)
        if verdict == "PASS":
            verdict_cell.font = pass_font
            verdict_cell.fill = pass_fill
        elif verdict:
            verdict_cell.font = fail_font
            verdict_cell.fill = fail_fill
        row += 1

    # 列宽
    widths = [6, 30, 12, 12, 12, 8, 10, 10, 20]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    wb.save(out_path)
    return os.path.abspath(out_path)
