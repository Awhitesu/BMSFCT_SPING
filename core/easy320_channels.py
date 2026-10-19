"""DO labels are octal; coil offsets remain ordinary zero-based integers."""
import re
from pathlib import Path


def normalize_channels(rows):
    result=[];seen=set()
    for row_number,label,name in rows:
        label=str(label or '').strip().upper()
        name='' if name is None else str(name).strip()
        if not label and not name:continue
        if not re.fullmatch(r'DO[0-7]+',label):
            raise ValueError(f'第{row_number}行序号 {label!r} 无效；请填写八进制编号，例如DO0、DO7、DO10')
        index=int(label[2:],8)
        if index>0xffff-0xfc08:raise ValueError(f'第{row_number}行编号超出设备地址范围')
        if index in seen:raise ValueError(f'第{row_number}行序号重复：DO{index:o}')
        seen.add(index)
        result.append({'do':f'DO{index:o}','name':name,'index':index})
    if not result:raise ValueError('表中没有有效DO通道')
    return sorted(result,key=lambda row:row['index'])


def read_channel_excel(path):
    from openpyxl import load_workbook
    if Path(path).suffix.lower() not in ('.xlsx','.xlsm'):
        raise ValueError('请将Excel另存为.xlsx格式后导入')
    book=load_workbook(path,read_only=True,data_only=False)
    try:
        sheet=book.active
        rows=sheet.iter_rows(values_only=True)
        headers=next(rows,())
        headers=[str(v or '').strip() for v in headers]
        if '序号' not in headers or '名称' not in headers:
            raise ValueError('第一行必须有“序号”和“名称”两列')
        serial,name=headers.index('序号'),headers.index('名称')
        return normalize_channels((n,row[serial] if len(row)>serial else None,
                                   row[name] if len(row)>name else None)
                                  for n,row in enumerate(rows,2))
    finally:book.close()


def default_channels():
    names={0:'平推气缸',1:'气缸抬起',2:'气缸下压'}
    return [{'do':f'DO{i:o}','name':names.get(i,''),'index':i} for i in range(32)]


def configured_channels(manager=None):
    database = getattr(manager, 'db_manager', None) if manager is not None else None
    if database is None:
        return default_channels()
    saved = (database.load_sys_config() or {}).get('easy320_channels')
    if saved is None:
        return default_channels()
    return normalize_channels((i+1, row['do'], row.get('name', '')) for i, row in enumerate(saved))
